"""Local MAPS graph, complete input accounting, and source-preserving extraction."""
from __future__ import annotations
import csv, hashlib, io, json, re, sqlite3, subprocess, sys, time, zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

HERE=Path(__file__).resolve().parent
LAB=HERE.parent/'maps-local-lab'
sys.path[:0]=[str(LAB),str(LAB/'.packages')]
import openpyxl
import prepare_data

def dump(path,obj):
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf8')

def decode(b):
    for enc in ('utf-8-sig','cp949','utf-16'):
        try:return b.decode(enc)
        except UnicodeError:pass
    return b.decode('utf8',errors='replace')

class Graph:
    def __init__(self): self.nodes={};self.edges={}
    def node(self,node_id,kind,label,**props):
        self.nodes[node_id]={'id':node_id,'kind':kind,'label':str(label),'properties':props};return node_id
    def edge(self,a,relation,b,**props):
        assert a in self.nodes and b in self.nodes,(a,b)
        self.edges[(a,relation,b)]={'source':a,'relation':relation,'target':b,'properties':props}

def build(backup=None,source=None):
    tick=time.perf_counter()
    old=json.loads((LAB/'data/snapshot.json').read_text(encoding='utf8'))
    backup=Path(backup or old['manifest'][0]['path'])
    source=Path(source) if source else next(Path(r['path']) for r in old['manifest'] if r['path'].endswith('consultation.csv')).parents[1]
    prepare_data.HERE=HERE
    data=prepare_data.prepare(backup,source)
    g=Graph();inventory=[];chunks=[];profiles=[];seen={};standards={}
    def ingest(path,b,parent=None,depth=0):
        sha=hashlib.sha256(b).hexdigest();id='doc:'+hashlib.sha256(path.encode()).hexdigest()[:20]
        ext=Path(path).suffix.lower()
        item={'path':path,'sha256':sha,'bytes':len(b),'extension':ext,'status':'','duplicate_of':'','error':''}
        inventory.append(item);g.node(id,'Document',Path(path).name,path=path,sha256=sha,bytes=len(b))
        if parent:g.edge(parent,'CONTAINS',id)
        if sha in seen:
            item.update(status='duplicate',duplicate_of=seen[sha]);g.edge(id,'DUPLICATE_OF',seen[sha]);return
        seen[sha]=id
        try:
            if ext in ('.zip','.dbc') and zipfile.is_zipfile(io.BytesIO(b)):
                if depth>=4: item['status']='archive_depth_limit';return
                with zipfile.ZipFile(io.BytesIO(b)) as z:
                    if sum(n.file_size for n in z.infolist())>500_000_000: item['status']='archive_size_limit';return
                    for n in z.infolist():
                        if not n.is_dir():ingest(path+'!'+n.filename,z.read(n),id,depth+1)
                item['status']='archive_expanded';return
            text=''
            if ext in ('.csv','.tsv'):
                rows=list(csv.DictReader(io.StringIO(decode(b)),delimiter='\t' if ext=='.tsv' else ','))
                heads=list(rows[0]) if rows else []
                profile={'document_id':id,'path':path,'rows':len(rows),'columns':len(heads),'headers':heads}
                profiles.append(profile)
                # Values from customer-level tables are deliberately not copied into the graph.
                for h in heads:
                    if h:
                        col=g.node('field:'+id+':'+h,'FileField',h)
                        g.edge(id,'HAS_FIELD',col)
                item['status']='structured_profiled';item['records']=len(rows)
                if '3.Standardization/' in path and '!' in path and path.startswith(str(backup)):
                    name=path.split('!')[-1]
                    standards[name]=rows
                return
            if ext in ('.xlsx','.xlsm'):
                wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
                for sh in wb:
                    profiles.append({'document_id':id,'path':path+'#'+sh.title,'rows':sh.max_row,'columns':sh.max_column,'headers':[]})
                    sheet=g.node('sheet:'+id+':'+sh.title,'Sheet',sh.title,rows=sh.max_row,columns=sh.max_column)
                    g.edge(id,'HAS_SHEET',sheet)
                wb.close();item['status']='structured_profiled';return
            if ext=='.json':
                obj=json.loads(decode(b));arr=obj if isinstance(obj,list) else next((v for v in obj.values() if isinstance(v,list)),[]) if isinstance(obj,dict) else []
                heads=list(arr[0]) if arr and isinstance(arr[0],dict) else list(obj) if isinstance(obj,dict) else []
                profiles.append({'document_id':id,'path':path,'rows':len(arr),'columns':len(heads),'headers':heads})
                item['status']='structured_profiled';item['records']=len(arr);return
            if ext in ('.docx','.pptx'):
                with zipfile.ZipFile(io.BytesIO(b)) as z:
                    names=[n for n in z.namelist() if n=='word/document.xml' or re.fullmatch(r'ppt/slides/slide\d+.xml',n)]
                    for n in names:
                        root=ET.fromstring(z.read(n));text+='\n'.join(t.text for t in root.iter() if t.tag.endswith('}t') and t.text)+'\n'
            elif ext=='.pdf':
                try:
                    from pypdf import PdfReader
                    text='\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(b)).pages)
                except ImportError:
                    runtime=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe'
                    proc=subprocess.run([str(runtime),str(HERE/'read_pdf.py')],input=b,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=60)
                    if proc.returncode:raise RuntimeError(proc.stderr.decode('utf8',errors='replace')[:500])
                    text=proc.stdout.decode('utf8')
            elif ext in ('.md','.txt','.dbml','.sql','.py','.ipynb','.html','.xml','.mf'):
                text=decode(b)
                if ext=='.html':text=re.sub('<[^>]+>',' ',text)
            else:item['status']='binary_inventory_only';return
            for offset in range(0,len(text),1600):
                chunks.append({'document_id':id,'offset':offset,'text':text[offset:offset+1600]})
            item['status']='document_indexed' if text.strip() else 'empty_text'
        except Exception as e:
            item.update(status='parse_error',error=type(e).__name__+': '+str(e))
    for p in sorted(source.rglob('*')):
        if p.is_file():ingest(str(p),p.read_bytes())
    ingest(str(backup),backup.read_bytes())
    def std(prefix):
        return next(v for k,v in standards.items() if Path(k).name.startswith(prefix))
    # Retain only the declared standard fields in catalog nodes; source examples are provenance only.
    for r in std('2.'):
        g.node('word:'+r['표준단어명'],'StandardWord',r['표준단어명'],abbr=r['표준단어약어'],english=r['영문FullName'],definition=r['정의'])
    for r in std('4.'):
        did=g.node('domain:'+r['도메인명'],'Domain',r['도메인명'],dtype=r['물리데이터유형'],definition=r['도메인설명'])
        rid=g.node('rule:'+r['도메인명'],'Rule',r['검증기준'],implementation='not_null_and_type_length_catalog_checks',raw=r['검증기준'])
        g.edge(did,'VALIDATED_BY',rid)
    term_rows=std('3.')
    for j,r in enumerate(term_rows):
        tid=g.node('term:'+str(j),'StandardTerm',r['표준용어명'],column=r['표준용어영문명_물리컬럼명'],english=r['논리영문FullName'],definition=r['정의'],entities=r['적용Entity'])
        if 'domain:'+r['도메인'] in g.nodes:g.edge(tid,'HAS_DOMAIN','domain:'+r['도메인'])
        for w in r['구성표준단어'].split('+'):
            if 'word:'+w in g.nodes:g.edge(tid,'COMPOSED_OF','word:'+w)
    for t in data['targets']:
        tab='target_table:'+t['table']
        if tab not in g.nodes:g.node(tab,'TargetTable',t['table'])
        tid=g.node('target:'+t['id'],'TargetColumn',t['id'],**t)
        g.edge(tab,'HAS_COLUMN',tid)
        for j,r in enumerate(term_rows):
            if t['column']==r['표준용어영문명_물리컬럼명'] and t['table'] in re.split(r'[,/\s]+',r['적용Entity']):g.edge(tid,'DEFINED_BY','term:'+str(j))
    for q in data['queries']+list(csv.DictReader((HERE/'data/unlabeled_columns.csv').open(encoding='utf-8-sig'))):
        qid=q.get('id') or '|'.join(q[k] for k in ('system','table','column'))
        sid='system:'+q['system'];tab='source_table:'+q['system']+'|'+q['table'].lower()
        if sid not in g.nodes:g.node(sid,'System',q['system'])
        if tab not in g.nodes:g.node(tab,'SourceTable',q['table']);g.edge(sid,'HAS_TABLE',tab)
        col=g.node('source:'+qid,'SourceColumn',q['column'],system=q['system'],table=q['table'],description=q['description'],dtype=q['dtype'],label_status='labeled' if 'gold' in q else 'unlabeled')
        g.edge(tab,'HAS_COLUMN',col)
        meta_names={'CHATBOT':'챗봇_Source_Metadata.xlsx','OUTBOUND':'tm_schema_code_product_v02.xlsx','HOMEPAGE':'homepage_complaint_metadata.csv','SALESFORCE':'salesforce_source_metadata_v1_2.csv','INBOUND':'# Inbound Source 특성 및 테스트 시나리오.md'}
        evidence=next((v for v in seen.values() if g.nodes[v]['properties']['path'].endswith(meta_names[q['system']])),None)
        if evidence:g.edge(col,'DESCRIBED_BY',evidence)
        for target in q.get('gold',[]):g.edge(col,'APPROVED_MAP','target:'+target,mapping_ids=q.get('mapping_ids',[]),usage='exploration_only; fold-specific benchmark graphs are rebuilt independently')
    for r in std('5.'):
        group='codeset:'+r['CODE_GROUP_ID']
        if group not in g.nodes:g.node(group,'CodeSet',r['CODE_GROUP_NAME'])
        code=g.node('code:'+r['CODE_GROUP_ID']+'|'+r['CODE_VALUE'],'Code',r['CODE_NAME'],value=r['CODE_VALUE'],valid=r['VALID_YN'],use=r['USE_YN'])
        g.edge(group,'HAS_VALUE',code)
    for t in data['targets']:
        if 'codeset:'+t['column'] in g.nodes:g.edge('target:'+t['id'],'USES_CODESET','codeset:'+t['column'])
    for j,r in enumerate(std('6.')):
        cid=g.node('codemap:'+str(j),'CodeMapping',r['SOURCE_COLUMN']+' = '+r['SOURCE_CODE'],**r)
        target='code:'+r['TARGET_CODE_GROUP_ID']+'|'+r['TARGET_CODE']
        if target in g.nodes:g.edge(cid,'CONVERTS_TO',target)
    # Every catalog entity has direct, inspectable evidence.
    catalog_doc=next(v for k,v in seen.items() if g.nodes[v]['properties']['path'].endswith('3.MAPS_표준용어사전_v1.0_확정.csv'))
    kind_prefix={'StandardWord':'2.','StandardTerm':'3.','Domain':'4.','CodeSet':'5.','Code':'5.','CodeMapping':'6.','TargetColumn':'8.','Rule':'4.'}
    for node in list(g.nodes.values()):
        if node['kind'] in kind_prefix:
            prefix=kind_prefix[node['kind']]
            doc=next((v for v in seen.values() if '3.Standardization/' in g.nodes[v]['properties']['path'] and Path(g.nodes[v]['properties']['path'].split('!')[-1]).name.startswith(prefix)),catalog_doc)
            g.edge(node['id'],'DERIVED_FROM',doc)
    out=HERE/'results';out.mkdir(exist_ok=True)
    db=sqlite3.connect(out/'knowledge.sqlite')
    db.executescript('DROP TABLE IF EXISTS nodes;DROP TABLE IF EXISTS edges;DROP TABLE IF EXISTS chunks;CREATE TABLE nodes(id TEXT PRIMARY KEY,kind TEXT,label TEXT,properties TEXT);CREATE TABLE edges(source TEXT,relation TEXT,target TEXT,properties TEXT,PRIMARY KEY(source,relation,target));CREATE INDEX edge_target ON edges(target,relation);CREATE TABLE chunks(document_id TEXT,offset INTEGER,text TEXT);')
    db.executemany('INSERT INTO nodes VALUES(:id,:kind,:label,:properties)',[{**n,'properties':json.dumps(n['properties'],ensure_ascii=False)} for n in g.nodes.values()])
    db.executemany('INSERT INTO edges VALUES(:source,:relation,:target,:properties)',[{**e,'properties':json.dumps(e['properties'],ensure_ascii=False)} for e in g.edges.values()])
    db.executemany('INSERT INTO chunks VALUES(:document_id,:offset,:text)',chunks)
    fts=True
    try:
        db.executescript('DROP TABLE IF EXISTS document_search;CREATE VIRTUAL TABLE document_search USING fts5(document_id UNINDEXED,text);INSERT INTO document_search SELECT document_id,text FROM chunks;')
    except sqlite3.OperationalError:fts=False
    db.commit();db.close()
    graph={'nodes':list(g.nodes.values()),'edges':list(g.edges.values())}
    dump(out/'graph.json',graph);dump(out/'inventory.json',inventory);dump(out/'profiles.json',profiles)
    # Provenance hashes are rechecked after extraction, independently of the cached snapshot.
    checks=[]
    for r in data['manifest']:
        path=r['path']
        if '!' in path:
            archive,member=path.split('!',1)
            with zipfile.ZipFile(archive) as z:b=z.read(member)
        else:b=Path(path).read_bytes()
        checks.append({'path':path,'unchanged':hashlib.sha256(b).hexdigest()==r['sha256']})
    assert all(c['unchanged'] for c in checks)
    summary={'physical_files':sum('!' not in r['path'] for r in inventory),'inventory_entries':len(inventory),'unique_contents':len(seen),'statuses':dict(Counter(r['status'] for r in inventory)),'nodes':len(g.nodes),'edges':len(g.edges),'node_types':dict(Counter(n['kind'] for n in g.nodes.values())),'edge_types':dict(Counter(e['relation'] for e in g.edges.values())),'document_chunks':len(chunks),'fts5_available':fts,'build_seconds':time.perf_counter()-tick,'database_bytes':(out/'knowledge.sqlite').stat().st_size,'source_checks':checks,'source_root':str(source),'backup':str(backup),'raw_customer_values_in_graph':False,'full_graph_contains_gold':'YES: exploration only. Evaluation uses independently constructed training-fold graphs.'}
    dump(out/'build_summary.json',summary)
    dump(HERE/'data/catalog.json',{'terms':term_rows,'domains':std('4.'),'words':std('2.'),'physical':std('8.'),'codes':std('5.'),'codemaps':std('6.')})
    print(json.dumps({k:v for k,v in summary.items() if k not in ['source_checks']},ensure_ascii=False,indent=2))
    return summary

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--backup');p.add_argument('--source-root');a=p.parse_args()
    build(a.backup,a.source_root)
