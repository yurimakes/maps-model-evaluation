"""Read original MAPS files without modifying them; create a metadata-only snapshot."""
from __future__ import annotations
import csv, hashlib, io, json, re, zipfile
from collections import defaultdict
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent

def csv_rows(data):
    for enc in ('utf-8-sig', 'cp949'):
        try:
            return list(csv.DictReader(io.StringIO(data.decode(enc))))
        except UnicodeDecodeError:
            pass
    raise ValueError('Unsupported CSV encoding')

def write_csv(path, rows):
    rows = list(rows)
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)

def prepare(backup, source_root):
    backup, source_root = Path(backup), Path(source_root)
    manifest = []
    def read_file(path):
        data = path.read_bytes()
        manifest.append({'path': str(path), 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()})
        return data
    with zipfile.ZipFile(io.BytesIO(read_file(backup))) as z:
        def member(suffix):
            names = [n for n in z.namelist() if n.endswith(suffix)]
            if len(names) != 1:
                raise ValueError(f'Expected unique archive member: {suffix}')
            data = z.read(names[0])
            manifest.append({'path': str(backup)+'!'+names[0], 'bytes':len(data), 'sha256':hashlib.sha256(data).hexdigest()})
            return data
        wb = openpyxl.load_workbook(io.BytesIO(member('MAPS_Source_Target_Mapping_v1.2.xlsx')), read_only=True, data_only=True)
        vals = list(wb['Mapping_Definition'].values)
        mappings = [dict(zip(vals[0], r)) for r in vals[1:] if r[0]]
        physical = csv_rows(member('8.MAPS_TO-BE_물리데이터모델_v1.0_확정.csv'))
        terms = csv_rows(member('3.MAPS_표준용어사전_v1.0_확정.csv'))
        words = csv_rows(member('2.MAPS_표준단어사전_v1.0_확정.csv'))
        old_runs = {}
        for n in ('self_consistency_raw.csv', 'self_consistency_raw_fixed.csv'):
            old_runs[n] = csv_rows(member(n))
    meta = {}
    def add(system, table, col, desc, dtype):
        meta[(system, str(table).lower(), str(col))] = {'system':system,'table':str(table),'column':str(col),'description':str(desc or ''),'dtype':str(dtype or '')}
    for system, pattern in [('HOMEPAGE','homepage_complaint_metadata.csv'), ('SALESFORCE','salesforce_source_metadata_v1_2.csv')]:
        paths = [p for p in source_root.rglob(pattern) if '이력' not in str(p)]
        if len(paths) != 1: raise ValueError(f'Expected source metadata: {pattern}')
        for r in csv_rows(read_file(paths[0])):
            table = 'homepage_complaint' if system == 'HOMEPAGE' else r['TABLE_NAME']
            add(system, table, r['COLUMN_NAME'], r['COLUMN_DESCRIPTION'], r['DATA_TYPE'])
    p = next(source_root.rglob('tm_schema_code_product_v02.xlsx'))
    w = openpyxl.load_workbook(io.BytesIO(read_file(p)),read_only=True,data_only=True)
    vals=list(w['컬럼 스키마'].values)
    for vals_row in vals[1:]:
        r=dict(zip(vals[0],vals_row))
        if r.get('컬럼명'):add('OUTBOUND','TM_SOURCE_EXPORT',r['컬럼명'],r['의미'],r['자료형'])
    p = next(source_root.rglob('챗봇_Source_Metadata.xlsx'))
    w = openpyxl.load_workbook(io.BytesIO(read_file(p)),read_only=True,data_only=True)
    vals=list(w.active.values); header=next(i for i,r in enumerate(vals) if '컬럼명' in r)
    for vals_row in vals[header+1:]:
        r=dict(zip(vals[header], vals_row))
        if r.get('컬럼명') and r.get('테이블명'):
            add('CHATBOT',r['테이블명'],r['컬럼명'],r['컬럼 설명'],r['데이터 타입'])
    p = next(source_root.rglob('# Inbound Source 특성 및 테스트 시나리오.md'))
    text=read_file(p).decode('utf-8-sig')
    for line in text.splitlines():
        cells=[c.strip() for c in line.strip().strip('|').split('|')]
        if len(cells)==8 and re.fullmatch(r'`[a-z_]+`',cells[1]):
            add('INBOUND','consultation',cells[1].strip('`'),cells[0],cells[2].strip('`'))
    gold=defaultdict(set); mapping_ids=defaultdict(list); pseudo=[]; unmapped_keys=[]
    for r in mappings:
        key=(r['SOURCE_SYSTEM'],r['SOURCE_TABLE'].lower(),r['SOURCE_COLUMN'])
        if str(r['SOURCE_COLUMN']).startswith('<'):
            pseudo.append({'system':r['SOURCE_SYSTEM'],'table':r['SOURCE_TABLE'],'column':r['SOURCE_COLUMN'],'target':f"{r['TARGET_TABLE']}.{r['TARGET_COLUMN']}"});continue
        mapping_ids[key].append(r['MAPPING_ID'])
        gold[key]
        if r['FINAL_MIGRATION_APPLY_YN']=='Y':
            if not r['TARGET_TABLE'] or not r['TARGET_COLUMN']:raise ValueError('Missing positive Target')
            gold[key].add(f"{r['TARGET_TABLE']}.{r['TARGET_COLUMN']}")
    queries=[]
    for key in sorted(gold):
        info=meta.get(key)
        if info is None:
            unmapped_keys.append('.'.join(key));info={'system':key[0],'table':key[1],'column':key[2],'description':'','dtype':''}
        queries.append({**info,'id':'|'.join(key),'gold':sorted(gold[key]),'mapping_ids':mapping_ids[key]})
    targets=[]
    for r in physical:
        matches=[t for t in terms if t['표준용어영문명_물리컬럼명']==r['COLUMN_NAME'] and r['TABLE_NAME'] in re.split(r'[,/\s]+',t['적용Entity'])]
        # Only standard definitions are used; SOURCE_근거 and mapping rules are deliberately excluded.
        term_text=' '.join(' '.join(t.get(k,'') for k in ['표준용어명','논리영문FullName','정의']) for t in matches)
        targets.append({'id':r['TABLE_NAME']+'.'+r['COLUMN_NAME'],'table':r['TABLE_NAME'],'column':r['COLUMN_NAME'],'description':r['DESCRIPTION'],'term':term_text,'dtype':r['DATA_TYPE']})
    ids={t['id'] for t in targets}
    assert len(ids)==len(targets)
    assert all(set(q['gold'])<=ids for q in queries)
    aliases=[]
    for r in words:
        aliases.append([r['표준단어약어'],r['영문FullName'],r['표준단어명']])
    coverage=[]
    for key,m in meta.items():
        if key not in gold:coverage.append({**m,'status':'UNLABELED_NOT_ASSUMED_NEGATIVE'})
    raw_inventory=[]
    source_files=['consultation.csv','tm_source_export_1000_v02.csv','homepage_complaint_source_1000_v1_2.csv','insurance_consultation_source_1000_v1_2.csv','chatbot_dummy_sessions_1000_dq.json']
    for name in source_files:
        paths=[p for p in source_root.rglob(name) if '이력' not in str(p)]
        if len(paths)!=1:continue
        p=paths[0];data=read_file(p)
        if p.suffix=='.csv':
            rows=csv_rows(data);raw_inventory.append({'file':str(p),'records':len(rows),'fields':len(rows[0]) if rows else 0})
        else:
            obj=json.loads(data.decode('utf-8-sig'));arr=obj if isinstance(obj,list) else next((v for v in obj.values() if isinstance(v,list)),[])
            raw_inventory.append({'file':str(p),'records':len(arr),'fields':len(arr[0]) if arr and isinstance(arr[0],dict) else 0})
    data={'queries':queries,'targets':targets,'aliases':aliases,'pseudo':pseudo,'old_runs':old_runs,'manifest':manifest,'raw_inventory':raw_inventory,'audit':{'mapping_rows':len(mappings),'positive_mapping_rows':sum(r['FINAL_MIGRATION_APPLY_YN']=='Y' for r in mappings),'missing_source_metadata':unmapped_keys,'unlabeled_metadata_columns':len(coverage),'excluded_pseudo_rows':len(pseudo)}}
    (HERE/'data').mkdir(exist_ok=True)
    (HERE/'data'/'snapshot.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    write_csv(HERE/'data'/'unlabeled_columns.csv',coverage)
    write_csv(HERE/'data'/'manifest.csv',manifest)
    write_csv(HERE/'data'/'queries.csv',[{**{k:v for k,v in q.items() if k not in ['gold','mapping_ids']},'gold':' ; '.join(q['gold'])} for q in queries])
    print(json.dumps(data['audit'],ensure_ascii=False,indent=2))
    print('queries',len(queries),'positive_queries',sum(bool(q['gold']) for q in queries),'gold_pairs',sum(len(q['gold']) for q in queries),'targets',len(targets))
    return data

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--backup',required=True);p.add_argument('--source-root',required=True);a=p.parse_args()
    prepare(a.backup,a.source_root)
